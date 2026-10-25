"""
converters.py - conversion logic for the Audio and Document Converter.

Nothing in this module imports Streamlit, so it can be tested or reused on its
own. Every error a user could reasonably cause is raised as ConversionError
with a message that is safe to show directly in the interface.
"""
from __future__ import annotations

import importlib.util
import io
import re
from dataclasses import dataclass
from functools import lru_cache
from typing import Callable, Optional

ProgressFn = Callable[[float], None]  # receives a fraction between 0 and 1

MAX_UPLOAD_MB = 50
TTS_CHUNK_CHARS = 4000       # gTTS handles long text best in pieces
GOOGLE_CHUNK_MS = 30_000     # Google's free recognizer rejects long clips


class ConversionError(Exception):
    """An error whose message is written for end users."""


# --------------------------------------------------------------------------- #
# Languages
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class Language:
    name: str
    tts: str       # gTTS language code
    whisper: str   # ISO 639-1 code used by Whisper
    google: str    # BCP-47 code used by Google Web Speech


LANGUAGES = [
    Language("English", "en", "en", "en-US"),
    Language("Hindi", "hi", "hi", "hi-IN"),
    Language("Bengali", "bn", "bn", "bn-IN"),
    Language("Spanish", "es", "es", "es-ES"),
    Language("French", "fr", "fr", "fr-FR"),
    Language("German", "de", "de", "de-DE"),
    Language("Italian", "it", "it", "it-IT"),
    Language("Portuguese", "pt", "pt", "pt-PT"),
    Language("Japanese", "ja", "ja", "ja-JP"),
    Language("Arabic", "ar", "ar", "ar-SA"),
]


def language_name(whisper_code: Optional[str]) -> str:
    """Turn a Whisper language code into a readable name."""
    for lang in LANGUAGES:
        if lang.whisper == whisper_code:
            return lang.name
    return whisper_code.upper() if whisper_code else "Unknown"


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def validate_upload(data: bytes, kind: str) -> None:
    """Reject empty, oversized or mislabelled files before doing any work."""
    if not data:
        raise ConversionError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise ConversionError(f"The file is larger than the {MAX_UPLOAD_MB} MB limit.")

    if kind == "pdf" and b"%PDF-" not in data[:1024]:
        raise ConversionError("This file is not a valid PDF. Check that it is not renamed or damaged.")

    if kind == "mp3":
        has_id3_tag = data[:3] == b"ID3"
        has_frame_sync = len(data) > 1 and data[0] == 0xFF and (data[1] & 0xE0) == 0xE0
        if not (has_id3_tag or has_frame_sync):
            raise ConversionError("This file is not a valid MP3. Check that it is not renamed or damaged.")


# --------------------------------------------------------------------------- #
# PDF to text
# --------------------------------------------------------------------------- #
def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_with_pdfplumber(data: bytes) -> str:
    """Second-chance extractor for PDFs that pypdf reads as blank."""
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return _clean("\n\n".join((page.extract_text() or "") for page in pdf.pages))
    except Exception:
        return ""


def extract_pdf_text(data: bytes, progress: Optional[ProgressFn] = None) -> tuple[str, int]:
    """Return (text, page_count). Raises ConversionError for unreadable PDFs."""
    validate_upload(data, "pdf")
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ConversionError("This PDF is password-protected. Remove the password and upload it again.")

        total = len(reader.pages)
        if total == 0:
            raise ConversionError("This PDF has no pages.")

        pages = []
        for index, page in enumerate(reader.pages):
            pages.append(page.extract_text() or "")
            if progress:
                progress((index + 1) / total)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError("This PDF appears to be corrupt or unreadable.") from exc

    text = _clean("\n\n".join(pages))
    if not text and importlib.util.find_spec("pdfplumber"):
        text = _extract_with_pdfplumber(data)
    return text, total


# --------------------------------------------------------------------------- #
# Text to MP3
# --------------------------------------------------------------------------- #
def split_text(text: str, limit: int = TTS_CHUNK_CHARS) -> list[str]:
    """Split text into chunks no longer than `limit`, preferring sentence ends."""
    sentences = re.split(r"(?<=[.!?\u0964\u061f])\s+", text.strip())
    chunks: list[str] = []
    current = ""

    for sentence in sentences:
        # A single sentence longer than the limit is wrapped at a space.
        while len(sentence) > limit:
            cut = sentence.rfind(" ", 0, limit)
            cut = cut if cut > 0 else limit
            if current:
                chunks.append(current)
                current = ""
            chunks.append(sentence[:cut])
            sentence = sentence[cut:].lstrip()

        if len(current) + len(sentence) + 1 > limit:
            if current:
                chunks.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()

    if current:
        chunks.append(current)
    return chunks


def text_to_mp3(text: str, lang_code: str = "en", slow: bool = False,
                progress: Optional[ProgressFn] = None) -> bytes:
    """Synthesize speech with gTTS and return the MP3 bytes."""
    text = text.strip()
    if not text:
        raise ConversionError("There is no text to convert.")

    from gtts import gTTS
    from gtts.tts import gTTSError

    chunks = split_text(text)
    output = io.BytesIO()
    for index, chunk in enumerate(chunks):
        try:
            gTTS(text=chunk, lang=lang_code, slow=slow).write_to_fp(output)
        except (gTTSError, AssertionError, ValueError) as exc:
            raise ConversionError(
                "Speech could not be generated. Check your internet connection, "
                "and that the text contains readable words in the chosen language."
            ) from exc
        if progress:
            progress((index + 1) / len(chunks))
    return output.getvalue()


# --------------------------------------------------------------------------- #
# MP3 to text
# --------------------------------------------------------------------------- #
@dataclass
class Transcript:
    text: str
    language: str
    duration: Optional[float]  # seconds
    engine: str


def whisper_available() -> bool:
    return importlib.util.find_spec("faster_whisper") is not None


@lru_cache(maxsize=2)
def _load_whisper(model_size: str):
    """Load (and cache for the life of the server) a faster-whisper model."""
    from faster_whisper import WhisperModel

    return WhisperModel(model_size, device="cpu", compute_type="int8")


def transcribe_whisper(data: bytes, model_size: str, language: Optional[Language],
                       progress: Optional[ProgressFn] = None) -> Transcript:
    """Transcribe locally with Whisper. `language=None` means auto-detect."""
    validate_upload(data, "mp3")
    try:
        model = _load_whisper(model_size)
    except Exception as exc:
        raise ConversionError(
            "The Whisper model could not be loaded. The first run needs an "
            "internet connection to download it."
        ) from exc

    try:
        segments, info = model.transcribe(
            io.BytesIO(data),
            language=language.whisper if language else None,
            vad_filter=True,  # skips silence, which also reduces hallucinated text
        )
        parts = []
        for segment in segments:  # lazy generator: decoding happens here
            parts.append(segment.text.strip())
            if progress and info.duration:
                progress(segment.end / info.duration)
    except Exception as exc:
        raise ConversionError("This audio could not be decoded. The file may be corrupt.") from exc

    text = " ".join(p for p in parts if p).strip()
    if not text:
        raise ConversionError("No speech was detected in this recording.")
    return Transcript(text, language_name(info.language), info.duration, "Whisper")


def transcribe_google(data: bytes, language: Language,
                      progress: Optional[ProgressFn] = None) -> Transcript:
    """Transcribe with Google's free web recognizer, in 30 second pieces."""
    validate_upload(data, "mp3")
    import speech_recognition as sr
    from pydub import AudioSegment

    try:
        audio = (
            AudioSegment.from_file(io.BytesIO(data), format="mp3")
            .set_channels(1)
            .set_frame_rate(16000)
        )
    except Exception as exc:
        raise ConversionError(
            "This audio could not be decoded. The file may be corrupt, "
            "or ffmpeg may not be installed on this machine."
        ) from exc

    recognizer = sr.Recognizer()
    parts = []
    total_ms = len(audio)
    for start in range(0, total_ms, GOOGLE_CHUNK_MS):
        buffer = io.BytesIO()
        audio[start:start + GOOGLE_CHUNK_MS].export(buffer, format="wav")
        buffer.seek(0)
        with sr.AudioFile(buffer) as source:
            clip = recognizer.record(source)
        try:
            parts.append(recognizer.recognize_google(clip, language=language.google))
        except sr.UnknownValueError:
            pass  # silent or unintelligible piece; keep going
        except sr.RequestError as exc:
            raise ConversionError(f"The Google speech service could not be reached: {exc}") from exc
        if progress:
            progress(min(start + GOOGLE_CHUNK_MS, total_ms) / total_ms)

    if not parts:
        raise ConversionError("No speech was detected in this recording.")
    return Transcript(" ".join(parts), language.name, total_ms / 1000, "Google Web Speech")


# --------------------------------------------------------------------------- #
# Exporters
# --------------------------------------------------------------------------- #
def text_to_docx(text: str, title: str = "Transcript") -> bytes:
    """Build a Word document from plain text and return its bytes."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ConversionError("Word export needs python-docx. Run: pip install python-docx") from exc

    document = Document()
    document.add_heading(title, level=1)
    for paragraph in text.splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
